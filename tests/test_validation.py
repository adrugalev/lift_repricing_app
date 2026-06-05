from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from src.calculator import calculate_project
from src.excel_io import export_repriced_pricing_excel, load_pricing_from_excel
from src.models import LiftInput, ProjectParams, TransferMethod
from src.validation import dataframe_to_lift_dicts, validate_lifts, validate_params


def test_lift_price_cannot_be_negative() -> None:
    with pytest.raises(ValidationError):
        LiftInput(
            lift_name="L1",
            original_lift_price_cny=-1,
            stops=1,
            installation_price_per_stop_rub=0,
        )


def test_stops_must_be_positive_integer() -> None:
    with pytest.raises(ValidationError):
        LiftInput(
            lift_name="L1",
            original_lift_price_cny=100,
            stops=0,
            installation_price_per_stop_rub=0,
        )


def test_installation_price_cannot_be_negative() -> None:
    with pytest.raises(ValidationError):
        LiftInput(
            lift_name="L1",
            original_lift_price_cny=100,
            stops=1,
            installation_price_per_stop_rub=-1,
        )


def test_exchange_rate_must_be_positive() -> None:
    params, errors = validate_params(
        {
            "exchange_rate_rub_per_cny": 0,
            "installation_markup": 0.1,
            "transfer_method": TransferMethod.PERCENT,
            "transfer_share": 0.2,
            "fixed_transfer_per_stop_rub": 0,
        }
    )

    assert params is None
    assert any("greater than 0" in error for error in errors)


def test_transfer_share_must_be_between_zero_and_one() -> None:
    with pytest.raises(ValidationError):
        ProjectParams(
            exchange_rate_rub_per_cny=10,
            installation_markup=0,
            transfer_method=TransferMethod.PERCENT,
            transfer_share=1.1,
            fixed_transfer_per_stop_rub=0,
        )


def test_currency_reserve_must_be_less_than_one() -> None:
    with pytest.raises(ValidationError):
        ProjectParams(
            exchange_rate_rub_per_cny=10,
            installation_markup=0,
            currency_reserve_share=1,
            transfer_method=TransferMethod.PERCENT,
            transfer_share=0.2,
            fixed_transfer_per_stop_rub=0,
        )


def test_excel_method_alias_sum_per_stop_is_supported() -> None:
    params = ProjectParams(
        exchange_rate_rub_per_cny=10,
        installation_markup=0,
        transfer_method="Сумма/ост.",
        transfer_share=0,
        fixed_transfer_per_stop_rub=50_000,
    )

    assert params.transfer_method == TransferMethod.FIXED


def test_dataframe_validation_reports_incomplete_rows() -> None:
    df = pd.DataFrame(
        [
            {
                "№": 1,
                "Лифт": "L1",
                "Цена лифта для клиента, CNY": 100_000,
                "Количество остановок": None,
                "Монтаж за 1 остановку без наценки, RUB": 200_000,
            }
        ]
    )

    lifts, errors = validate_lifts(dataframe_to_lift_dicts(df))

    assert lifts == []
    assert errors


def test_load_pricing_file_from_cost_sheet(tmp_path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Стоимость"
    ws["O1"] = "Маржа на оборудование"
    ws["P1"] = 0.14
    ws["Q1"] = "Маржа на  МОНТАЖ"
    ws["R1"] = 0.05
    ws["A2"] = "Продукт"
    ws["B2"] = "Кол-во лифтов"
    ws["C2"] = "Кол-во этажей"
    ws["I2"] = "Общая стоимость для клиента"
    ws["J2"] = "Стоимость монтажа за этаж"
    ws["Q2"] = "Наценка МОНТАЖ"
    ws["R2"] = 0.052631578947368425
    ws["A3"] = "L1"
    ws["B3"] = 1
    ws["C3"] = 9
    ws["I3"] = 180_200
    ws["J3"] = 215_000
    ws["A4"] = "L2"
    ws["B4"] = 2
    ws["C4"] = 8
    ws["I4"] = 153_000
    ws["J4"] = 215_000
    ws["A5"] = "Итого"
    ws["A32"] = "Курс рубля у Юаню"
    ws["E32"] = 11.25
    path = tmp_path / "pricing.xlsx"
    workbook.save(path)

    df, params, warnings = load_pricing_from_excel(path)

    assert warnings == []
    assert len(df) == 3
    assert list(df["Лифт"]) == ["L1", "L2 #1", "L2 #2"]
    assert df["Цена лифта для клиента, CNY"].sum() == 486_200
    assert params is not None
    assert params.exchange_rate_rub_per_cny == 11.25
    assert params.installation_markup == pytest.approx(0.052631578947368425)


def test_export_repriced_pricing_file_updates_original_cost_sheet(tmp_path) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Стоимость"
    ws["A2"] = "Продукт"
    ws["B2"] = "Кол-во лифтов"
    ws["C2"] = "Кол-во этажей"
    ws["I2"] = "Общая стоимость для клиента"
    ws["J2"] = "Стоимость монтажа за этаж"
    ws["K2"] = "Общая стоимость монтажа"
    ws["L2"] = "Наценка на монтаже"
    ws["M2"] = "Общая стоимость монтажа для клиента"
    ws["Q2"] = "Наценка МОНТАЖ"
    ws["R2"] = 0.1
    ws["A3"] = "L1"
    ws["B3"] = 1
    ws["C3"] = 2
    ws["I3"] = 100_000
    ws["J3"] = 100_000
    ws["K3"] = "=J3*C3"
    ws["L3"] = "=K3*$R$2"
    ws["M3"] = "=K3+(K3*$R$2)"
    ws["A4"] = "Итого"
    ws["I4"] = "=SUM(I3:I3)"
    ws["K4"] = "=SUM(K3:K3)"
    ws["L4"] = "=SUM(L3:L3)"
    ws["M4"] = "=K4+(K4*$R$2)"
    path = tmp_path / "pricing.xlsx"
    workbook.save(path)

    params = ProjectParams(
        exchange_rate_rub_per_cny=10,
        installation_markup=0.1,
        transfer_method=TransferMethod.PERCENT,
        transfer_share=0.2,
        fixed_transfer_per_stop_rub=0,
    )
    lifts = [
        LiftInput(
            number=1,
            lift_name="L1",
            original_lift_price_cny=100_000,
            stops=2,
            installation_price_per_stop_rub=100_000,
        )
    ]
    results, _ = calculate_project(lifts, params)

    output = export_repriced_pricing_excel(path, results, params)
    updated = load_workbook(BytesIO(output.getvalue()), data_only=False)
    updated_ws = updated["Стоимость"]

    assert isinstance(updated_ws["I3"].value, str)
    assert updated_ws["I3"].value.startswith("=ROUND(")
    assert "$R$2" not in updated_ws["I3"].value
    assert isinstance(updated_ws["J3"].value, str)
    assert updated_ws["J3"].value.startswith("=ROUND(")
    assert "1-0.2" in updated_ws["J3"].value
    assert updated_ws["K3"].value == "=J3*C3"
    assert updated_ws["L3"].value == "=K3*$R$2"
    assert updated_ws["M3"].value == "=ROUNDUP((K3+(K3*$R$2)),-3)"
    assert updated_ws["I4"].value == "=SUM(I3:I3)"
    assert updated_ws["K4"].value == "=SUM(K3:K3)"
    assert updated_ws["L4"].value == "=SUM(L3:L3)"
    assert updated_ws["M4"].value == "=ROUNDUP((K4+(K4*$R$2)),-3)"

    with ZipFile(BytesIO(output.getvalue())) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert "<f>ROUND(" in sheet_xml
        assert "<v>104400</v>" in sheet_xml
        assert "<v>80000</v>" in sheet_xml
        assert "<v>160000</v>" in sheet_xml
        assert "<v>16000</v>" in sheet_xml
        assert "<v>176000</v>" in sheet_xml
