from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO
import math
from pathlib import Path
import re
from typing import Any, BinaryIO
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape as xml_escape
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .formatter import CNY_NUMBER_FORMAT, PERCENT_NUMBER_FORMAT, RUB_NUMBER_FORMAT
from .models import LiftInput, LiftResult, ProjectParams, ProjectSummary, TransferMethod
from .validation import INPUT_COLUMNS, is_blank

GROUP_SHEET_NAME = "Групповая переоценка"
CALCULATOR_SHEET_NAME = "Калькулятор"
PRICING_SHEET_NAME = "Стоимость"

RESULT_COLUMNS = [
    "№",
    "Лифт",
    "Исходная цена лифта, CNY",
    "Количество остановок",
    "Монтаж за 1 остановку без наценки, RUB",
    "Монтаж всего до переноса, RUB",
    "Монтаж за 1 остановку с наценкой, RUB",
    "Перенос за 1 остановку, RUB",
    "Сумма переноса, RUB",
    "Перенесено в лифт, CNY",
    "Новая цена лифта, CNY",
    "Новый монтаж, RUB",
    "Новый монтаж за 1 остановку, RUB",
    "Проверка / статус",
]


def _cell_value(ws: Any, coord: str) -> Any:
    return ws[coord].value if ws is not None else None


def load_pricing_from_excel(source: str | BinaryIO | BytesIO) -> tuple[pd.DataFrame, ProjectParams | None, list[str]]:
    workbook = load_workbook(source, data_only=True)
    ws, header_row, columns = _find_pricing_table(workbook)

    rows: list[dict[str, Any]] = []
    warnings: list[str] = []
    number = 1
    lift_rows_count = 0
    rows_without_installation = 0

    for row_idx in range(header_row + 1, ws.max_row + 1):
        lift_name = ws.cell(row_idx, columns["lift_name"]).value
        if is_blank(lift_name):
            continue
        if str(lift_name).strip().lower() == "итого":
            break

        lift_rows_count += 1
        qty = _positive_int_or_one(ws.cell(row_idx, columns["quantity"]).value)
        price_cny = ws.cell(row_idx, columns["price_cny"]).value
        stops = ws.cell(row_idx, columns["stops"]).value
        installation_per_stop = ws.cell(row_idx, columns["installation_per_stop"]).value

        if is_blank(installation_per_stop):
            rows_without_installation += 1

        if any(is_blank(value) for value in [price_cny, stops, installation_per_stop]):
            warnings.append(f"Строка {row_idx}: пропущена, не заполнены цена, остановки или монтаж.")
            continue

        for copy_idx in range(qty):
            display_name = str(lift_name).strip()
            if qty > 1:
                display_name = f"{display_name} #{copy_idx + 1}"
            rows.append(
                {
                    "№": number,
                    "Лифт": display_name,
                    "Цена лифта для клиента, CNY": price_cny,
                    "Количество остановок": stops,
                    "Монтаж за 1 остановку без наценки, RUB": installation_per_stop,
                }
            )
            number += 1

    if lift_rows_count > 0 and not rows and rows_without_installation == lift_rows_count:
        raise ValueError(
            "В этой расценке нет стоимости монтажа. Соответственно этот файл невозможно "
            "обработать для групповой переоценки, потому что перенос выполняется из "
            "рублевой стоимости монтажа в стоимость лифта."
        )

    params = _extract_pricing_params(ws)
    if params is None:
        warnings.append("Курс или наценка монтажа из файла расценки не найдены, оставлены текущие параметры.")

    return pd.DataFrame(rows, columns=INPUT_COLUMNS), params, warnings


def load_from_excel(source: str | BinaryIO | BytesIO) -> tuple[pd.DataFrame, ProjectParams | None, list[str]]:
    workbook = load_workbook(source, data_only=True)
    group_ws = workbook[GROUP_SHEET_NAME] if GROUP_SHEET_NAME in workbook.sheetnames else workbook.worksheets[1]
    calc_ws = workbook[CALCULATOR_SHEET_NAME] if CALCULATOR_SHEET_NAME in workbook.sheetnames else None

    rows: list[dict[str, Any]] = []
    for row_idx in range(11, group_ws.max_row + 1):
        raw = {
            "№": _cell_value(group_ws, f"B{row_idx}"),
            "Лифт": _cell_value(group_ws, f"C{row_idx}"),
            "Цена лифта для клиента, CNY": _cell_value(group_ws, f"D{row_idx}"),
            "Количество остановок": _cell_value(group_ws, f"E{row_idx}"),
            "Монтаж за 1 остановку без наценки, RUB": _cell_value(group_ws, f"F{row_idx}"),
        }
        if all(
            is_blank(raw[column])
            for column in [
                "Лифт",
                "Цена лифта для клиента, CNY",
                "Монтаж за 1 остановку без наценки, RUB",
            ]
        ):
            continue
        rows.append(raw)

    params: ProjectParams | None = None
    warnings: list[str] = []
    if calc_ws is not None:
        try:
            params = ProjectParams(
                exchange_rate_rub_per_cny=float(_cell_value(calc_ws, "C7") or 0),
                installation_markup=float(_cell_value(calc_ws, "C13") or 0),
                transfer_method=TransferMethod.from_raw(_cell_value(calc_ws, "C10") or TransferMethod.PERCENT),
                transfer_share=float(_cell_value(calc_ws, "C11") or 0),
                fixed_transfer_per_stop_rub=float(_cell_value(calc_ws, "C12") or 0),
            )
        except Exception as exc:
            warnings.append(f"Параметры с листа '{CALCULATOR_SHEET_NAME}' не загружены: {exc}")
    else:
        warnings.append(f"Лист '{CALCULATOR_SHEET_NAME}' не найден, параметры нужно заполнить вручную.")

    return pd.DataFrame(rows, columns=INPUT_COLUMNS), params, warnings


def _find_pricing_table(workbook: Any) -> tuple[Any, int, dict[str, int]]:
    candidates = []
    if PRICING_SHEET_NAME in workbook.sheetnames:
        candidates.append(workbook[PRICING_SHEET_NAME])
    candidates.extend(ws for ws in workbook.worksheets if ws.title != PRICING_SHEET_NAME)

    for ws in candidates:
        for row_idx in range(1, min(ws.max_row, 20) + 1):
            header_values = {
                _normalize_header(ws.cell(row_idx, col_idx).value): col_idx
                for col_idx in range(1, ws.max_column + 1)
                if not is_blank(ws.cell(row_idx, col_idx).value)
            }
            columns = _match_pricing_columns(header_values)
            if columns is not None:
                return ws, row_idx, columns

    raise ValueError(
        "Не найдена таблица расценки. Нужны заголовки: Продукт, Кол-во этажей, "
        "Общая стоимость для клиента, Стоимость монтажа за этаж."
    )


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("\n", " ").split())


def _match_pricing_columns(headers: dict[str, int]) -> dict[str, int] | None:
    def find(*needles: str) -> int | None:
        normalized_needles = [_normalize_header(needle) for needle in needles]
        for header, col_idx in headers.items():
            if any(needle in header for needle in normalized_needles):
                return col_idx
        return None

    lift_name = find("продукт")
    quantity = find("кол-во лифтов", "количество лифтов")
    stops = find("кол-во этажей", "количество этажей", "остановки")
    price_cny = find("общая стоимость для клиента", "цена лифта для клиента")
    installation_per_stop = find(
        "стоимость монтажа за этаж",
        "монтаж за 1 остановку",
        "себестоимость монтажа",
        "монтаж/демонтаж",
    )

    if all([lift_name, stops, price_cny]) and not installation_per_stop:
        raise ValueError(
            "В этой расценке нет колонки со стоимостью монтажа. Соответственно этот файл "
            "невозможно обработать для групповой переоценки, потому что перенос выполняется "
            "из рублевой стоимости монтажа в стоимость лифта."
        )

    if not all([lift_name, stops, price_cny, installation_per_stop]):
        return None

    return {
        "lift_name": lift_name,
        "quantity": quantity or lift_name,
        "stops": stops,
        "price_cny": price_cny,
        "installation_per_stop": installation_per_stop,
    }


def _positive_int_or_one(value: Any) -> int:
    try:
        numeric = int(float(value))
    except (TypeError, ValueError):
        return 1
    return numeric if numeric > 0 else 1


def _extract_pricing_params(ws: Any) -> ProjectParams | None:
    exchange_rate = _find_value_by_label(ws, "курс рубля")
    installation_markup = _find_value_by_label(ws, "наценка монтаж", preferred_col_offset=1)
    if installation_markup is None:
        installation_markup = _find_value_by_label(ws, "маржа на монтаж", preferred_col_offset=1)
        if installation_markup is not None:
            installation_markup = installation_markup / (1 - installation_markup)

    if exchange_rate is None or installation_markup is None:
        return None

    try:
        return ProjectParams(
            exchange_rate_rub_per_cny=float(exchange_rate),
            installation_markup=float(installation_markup),
            transfer_method=TransferMethod.PERCENT,
            transfer_share=0.2,
            fixed_transfer_per_stop_rub=0.0,
        )
    except Exception:
        return None


def _find_value_by_label(ws: Any, label_fragment: str, preferred_col_offset: int = 4) -> float | None:
    needle = _normalize_header(label_fragment)
    for row in ws.iter_rows():
        for cell in row:
            if needle in _normalize_header(cell.value):
                preferred = ws.cell(cell.row, cell.column + preferred_col_offset).value
                numeric = _try_float(preferred)
                if numeric is not None:
                    return numeric
                for col_idx in range(cell.column + 1, min(ws.max_column, cell.column + 8) + 1):
                    numeric = _try_float(ws.cell(cell.row, col_idx).value)
                    if numeric is not None:
                        return numeric
    return None


def _try_float(value: Any) -> float | None:
    if is_blank(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def lifts_to_dataframe(lifts: Iterable[LiftInput]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "№": lift.number,
                "Лифт": lift.lift_name,
                "Цена лифта для клиента, CNY": lift.original_lift_price_cny,
                "Количество остановок": lift.stops,
                "Монтаж за 1 остановку без наценки, RUB": lift.installation_price_per_stop_rub,
            }
            for lift in lifts
        ],
        columns=INPUT_COLUMNS,
    )


def results_to_dataframe(results: Iterable[LiftResult]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "№": row.number,
                "Лифт": row.lift_name,
                "Исходная цена лифта, CNY": row.original_lift_price_cny,
                "Количество остановок": row.stops,
                "Монтаж за 1 остановку без наценки, RUB": row.installation_price_per_stop_rub,
                "Монтаж всего до переноса, RUB": row.original_installation_total_rub,
                "Монтаж за 1 остановку с наценкой, RUB": row.installation_per_stop_with_markup_rub,
                "Перенос за 1 остановку, RUB": row.transfer_per_stop_rub,
                "Сумма переноса, RUB": row.transfer_total_rub,
                "Перенесено в лифт, CNY": row.transferred_to_lift_cny,
                "Новая цена лифта, CNY": row.new_lift_price_cny,
                "Новый монтаж, RUB": row.new_installation_total_rub,
                "Новый монтаж за 1 остановку, RUB": row.new_installation_per_stop_rub,
                "Проверка / статус": row.status,
            }
            for row in results
        ],
        columns=RESULT_COLUMNS,
    )


def summary_to_dataframe(summary: ProjectSummary, params: ProjectParams | None = None) -> pd.DataFrame:
    rows = [
        *(
            [
                ("Курс из расценки, RUB за 1 CNY", params.exchange_rate_rub_per_cny, "RATE"),
                ("Валютный резерв", params.currency_reserve_share, "PERCENT"),
                ("Курс переноса, RUB за 1 CNY", params.transfer_exchange_rate_rub_per_cny, "RATE"),
            ]
            if params is not None
            else []
        ),
        ("Общая исходная стоимость лифтов, CNY", summary.total_original_lifts_cny, "CNY"),
        ("Общая новая стоимость лифтов, CNY", summary.total_new_lifts_cny, "CNY"),
        ("Разница по лифтам, CNY", summary.lift_price_delta_cny, "CNY"),
        ("Общая исходная стоимость монтажа, RUB", summary.total_original_installation_rub, "RUB"),
        ("Общая новая стоимость монтажа, RUB", summary.total_new_installation_rub, "RUB"),
        ("Перенесено из монтажа всего, RUB", summary.total_transferred_rub, "RUB"),
        ("Перенесено в лифты всего, CNY", summary.total_transferred_cny, "CNY"),
        ("Итог проекта до переноса, CNY", summary.total_project_before_cny, "CNY"),
        ("Итог проекта после переноса, CNY", summary.total_project_after_cny, "CNY"),
        ("Контроль проекта, CNY", summary.project_control_cny, "CNY"),
        ("Статус", summary.status, "TEXT"),
    ]
    return pd.DataFrame(rows, columns=["Показатель", "Значение", "Валюта"])


def export_results_to_excel(
    inputs_df: pd.DataFrame,
    results: list[LiftResult],
    summary: ProjectSummary,
    params: ProjectParams,
) -> BytesIO:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    _write_inputs_sheet(workbook, inputs_df, params)
    _write_results_sheet(workbook, results_to_dataframe(results))
    _write_summary_sheet(workbook, summary_to_dataframe(summary, params))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def export_repriced_pricing_excel(
    source: str | BinaryIO | BytesIO,
    results: list[LiftResult],
    params: ProjectParams,
) -> BytesIO:
    source_bytes = _read_source_bytes(source)
    workbook = load_workbook(BytesIO(source_bytes), data_only=False)
    ws, header_row, columns = _find_pricing_table(workbook)
    exchange_rate_cell = _find_value_cell_by_label(ws, "курс рубля", preferred_col_offset=4)
    markup_cell = _find_value_cell_by_label(ws, "наценка монтаж", preferred_col_offset=1)
    markup_formula = _absolute_ref(markup_cell) if markup_cell is not None else str(params.installation_markup)
    exchange_formula = (
        _absolute_ref(exchange_rate_cell)
        if exchange_rate_cell is not None
        else str(params.exchange_rate_rub_per_cny)
    )

    formulas: dict[str, tuple[str, float]] = {}
    result_idx = 0
    data_row_indices: list[int] = []
    cached_price_values: list[float] = []
    cached_installation_total_values: list[float] = []
    cached_installation_markup_values: list[float] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        lift_name = ws.cell(row_idx, columns["lift_name"]).value
        if is_blank(lift_name):
            continue
        if str(lift_name).strip().lower() == "итого":
            total_row_idx = row_idx
            break

        qty = _positive_int_or_one(ws.cell(row_idx, columns["quantity"]).value)
        row_results = results[result_idx : result_idx + qty]
        if not row_results:
            break

        old_price = _formula_operand(ws.cell(row_idx, columns["price_cny"]))
        old_installation_per_stop = _formula_operand(
            ws.cell(row_idx, columns["installation_per_stop"])
        )

        if params.transfer_method == TransferMethod.PERCENT:
            transfer_per_stop = (
                f"({old_installation_per_stop})*(1+{markup_formula})*"
                f"{_excel_number(params.transfer_share)}"
            )
            new_installation_per_stop = (
                f"({old_installation_per_stop})*(1-{_excel_number(params.transfer_share)})"
            )
        else:
            transfer_per_stop = (
                f"MIN({_excel_number(params.fixed_transfer_per_stop_rub)},"
                f"({old_installation_per_stop})*(1+{markup_formula}))"
            )
            new_installation_per_stop = (
                f"MAX(0,({old_installation_per_stop})-({transfer_per_stop})/(1+{markup_formula}))"
            )

        cached_price = round(sum(row.new_lift_price_cny for row in row_results) / len(row_results))
        cached_transfer_to_lift_cny = (
            sum(row.transferred_to_lift_cny for row in row_results) / len(row_results)
        )
        cached_base_per_stop = round(
            (
                sum(row.new_installation_total_rub / row.stops for row in row_results)
                / len(row_results)
                / (1 + params.installation_markup)
            )
        )
        cached_installation_total = cached_base_per_stop * int(float(ws.cell(row_idx, columns["stops"]).value or 0))
        cached_installation_markup = cached_installation_total * params.installation_markup
        cached_installation_client_total = _roundup(cached_installation_total + cached_installation_markup, -3)

        formulas[ws.cell(row_idx, columns["price_cny"]).coordinate] = (
            f"ROUND(({old_price})+{_excel_number(cached_transfer_to_lift_cny)},0)",
            cached_price,
        )
        formulas[ws.cell(row_idx, columns["installation_per_stop"]).coordinate] = (
            f"ROUND({new_installation_per_stop},0)",
            cached_base_per_stop,
        )
        _add_pricing_installation_formulas(
            ws,
            columns,
            header_row,
            row_idx,
            markup_formula,
            cached_installation_total,
            cached_installation_markup,
            cached_installation_client_total,
            formulas,
        )

        data_row_indices.append(row_idx)
        cached_price_values.append(cached_price)
        cached_installation_total_values.append(cached_installation_total)
        cached_installation_markup_values.append(cached_installation_markup)

        result_idx += qty

    if "total_row_idx" in locals() and data_row_indices:
        _add_pricing_total_formulas(
            ws,
            columns,
            header_row,
            total_row_idx,
            data_row_indices,
            markup_formula,
            sum(cached_price_values),
            sum(cached_installation_total_values),
            sum(cached_installation_markup_values),
            formulas,
        )

    return _patch_xlsx_formulas(source_bytes, ws.title, formulas)


def _read_source_bytes(source: str | BinaryIO | BytesIO) -> bytes:
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    position = source.tell()
    source.seek(0)
    data = source.read()
    source.seek(position)
    return data


def _excel_number(value: float) -> str:
    return f"{float(value):.12g}"


def _roundup(value: float, digits: int = 0) -> float:
    factor = 10 ** (-digits)
    return math.ceil(value / factor) * factor


def _add_pricing_installation_formulas(
    ws: Any,
    columns: dict[str, int],
    header_row: int,
    row_idx: int,
    markup_formula: str,
    cached_installation_total: float,
    cached_installation_markup: float,
    cached_installation_client_total: float,
    formulas: dict[str, tuple[str, float]],
) -> None:
    installation_total_col = _find_column_by_header(
        ws,
        header_row,
        columns,
        "общая стоимость монтажа",
        exclude=("для клиента",),
    )
    installation_markup_col = _find_column_by_header(ws, header_row, columns, "наценка на монтаже")
    installation_client_total_col = _find_column_by_header(
        ws,
        header_row,
        columns,
        "общая стоимость монтажа для клиента",
    )

    stops_coord = ws.cell(row_idx, columns["stops"]).coordinate
    installation_per_stop_coord = ws.cell(row_idx, columns["installation_per_stop"]).coordinate

    if installation_total_col is not None:
        total_coord = ws.cell(row_idx, installation_total_col).coordinate
        formulas[total_coord] = (
            f"{installation_per_stop_coord}*{stops_coord}",
            cached_installation_total,
        )
    else:
        total_coord = ""

    if installation_markup_col is not None and total_coord:
        markup_coord = ws.cell(row_idx, installation_markup_col).coordinate
        formulas[markup_coord] = (f"{total_coord}*{markup_formula}", cached_installation_markup)

    if installation_client_total_col is not None and total_coord:
        client_total_coord = ws.cell(row_idx, installation_client_total_col).coordinate
        formulas[client_total_coord] = (
            f"ROUNDUP(({total_coord}+({total_coord}*{markup_formula})),-3)",
            cached_installation_client_total,
        )
        next_cell = ws.cell(row_idx, installation_client_total_col + 1)
        if next_cell.value is not None and not is_blank(ws.cell(row_idx, columns["stops"]).value):
            formulas[next_cell.coordinate] = (
                f"{client_total_coord}/{stops_coord}",
                cached_installation_client_total / float(ws.cell(row_idx, columns["stops"]).value),
            )


def _add_pricing_total_formulas(
    ws: Any,
    columns: dict[str, int],
    header_row: int,
    total_row_idx: int,
    data_row_indices: list[int],
    markup_formula: str,
    cached_price_total: float,
    cached_installation_total: float,
    cached_installation_markup: float,
    formulas: dict[str, tuple[str, float]],
) -> None:
    first_row = data_row_indices[0]
    last_row = data_row_indices[-1]
    installation_total_col = _find_column_by_header(
        ws,
        header_row,
        columns,
        "общая стоимость монтажа",
        exclude=("для клиента",),
    )
    installation_markup_col = _find_column_by_header(ws, header_row, columns, "наценка на монтаже")
    installation_client_total_col = _find_column_by_header(
        ws,
        header_row,
        columns,
        "общая стоимость монтажа для клиента",
    )

    price_letter = get_column_letter(columns["price_cny"])
    formulas[ws.cell(total_row_idx, columns["price_cny"]).coordinate] = (
        f"SUM({price_letter}{first_row}:{price_letter}{last_row})",
        cached_price_total,
    )

    if installation_total_col is not None:
        total_letter = get_column_letter(installation_total_col)
        total_coord = ws.cell(total_row_idx, installation_total_col).coordinate
        formulas[total_coord] = (
            f"SUM({total_letter}{first_row}:{total_letter}{last_row})",
            cached_installation_total,
        )
    else:
        total_coord = ""

    if installation_markup_col is not None:
        markup_letter = get_column_letter(installation_markup_col)
        formulas[ws.cell(total_row_idx, installation_markup_col).coordinate] = (
            f"SUM({markup_letter}{first_row}:{markup_letter}{last_row})",
            cached_installation_markup,
        )

    if installation_client_total_col is not None and total_coord:
        formulas[ws.cell(total_row_idx, installation_client_total_col).coordinate] = (
            f"ROUNDUP(({total_coord}+({total_coord}*{markup_formula})),-3)",
            _roundup(cached_installation_total + cached_installation_markup, -3),
        )


def _find_column_by_header(
    ws: Any,
    header_row: int,
    columns: dict[str, int],
    include: str,
    exclude: tuple[str, ...] = (),
) -> int | None:
    include_normalized = _normalize_header(include)
    exclude_normalized = tuple(_normalize_header(value) for value in exclude)
    for col_idx in range(1, ws.max_column + 1):
        header = _normalize_header(ws.cell(header_row, col_idx).value)
        if include_normalized in header and not any(value in header for value in exclude_normalized):
            return col_idx
    return None


def _formula_operand(cell: Any) -> str:
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return f"({value[1:]})"
    return _excel_number(float(value or 0))


def _absolute_ref(cell: Any) -> str:
    return f"${get_column_letter(cell.column)}${cell.row}"


def _find_value_cell_by_label(ws: Any, label_fragment: str, preferred_col_offset: int = 4) -> Any | None:
    needle = _normalize_header(label_fragment)
    for row in ws.iter_rows():
        for cell in row:
            if needle in _normalize_header(cell.value):
                preferred = ws.cell(cell.row, cell.column + preferred_col_offset)
                if _try_float(preferred.value) is not None or (
                    isinstance(preferred.value, str) and preferred.value.startswith("=")
                ):
                    return preferred
                for col_idx in range(cell.column + 1, min(ws.max_column, cell.column + 8) + 1):
                    candidate = ws.cell(cell.row, col_idx)
                    if _try_float(candidate.value) is not None or (
                        isinstance(candidate.value, str) and candidate.value.startswith("=")
                    ):
                        return candidate
    return None


def _patch_xlsx_formulas(
    source_bytes: bytes,
    sheet_name: str,
    formulas: dict[str, tuple[str, float]],
) -> BytesIO:
    sheet_path = _sheet_xml_path(source_bytes, sheet_name)
    output = BytesIO()
    with ZipFile(BytesIO(source_bytes), "r") as zin:
        with ZipFile(output, "w", ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "xl/calcChain.xml":
                    continue
                data = zin.read(item.filename)
                if item.filename == sheet_path:
                    data = _patch_sheet_xml(data, formulas)
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    data = _remove_calc_chain_relationship(data)
                elif item.filename == "[Content_Types].xml":
                    data = _remove_calc_chain_content_type(data)
                zout.writestr(item, data)
    output.seek(0)
    return output


def _sheet_xml_path(source_bytes: bytes, sheet_name: str) -> str:
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    with ZipFile(BytesIO(source_bytes), "r") as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        rel_id = None
        for sheet in workbook_root.findall("main:sheets/main:sheet", ns):
            if sheet.attrib.get("name") == sheet_name:
                rel_id = sheet.attrib[f"{{{ns['rel']}}}id"]
                break
        if rel_id is None:
            raise ValueError(f"Лист '{sheet_name}' не найден в книге.")

        rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        for relationship in rels_root.findall("pkgrel:Relationship", ns):
            if relationship.attrib.get("Id") == rel_id:
                target = relationship.attrib["Target"].lstrip("/")
                return target if target.startswith("xl/") else f"xl/{target}"
    raise ValueError(f"XML-файл листа '{sheet_name}' не найден.")


def _patch_sheet_xml(xml_bytes: bytes, formulas: dict[str, tuple[str, float]]) -> bytes:
    xml = xml_bytes.decode("utf-8")
    for coord, (formula, cached_value) in formulas.items():
        pattern = re.compile(rf'(<c\b[^>]*\br="{re.escape(coord)}"[^>]*>)(.*?</c>)', re.DOTALL)
        match = pattern.search(xml)
        if match is None:
            raise ValueError(f"Ячейка {coord} не найдена в XML листа.")
        open_tag = re.sub(r'\s+t="[^"]*"', "", match.group(1))
        replacement = (
            f"{open_tag}<f>{xml_escape(formula)}</f>"
            f"<v>{_excel_number(cached_value)}</v></c>"
        )
        xml = xml[: match.start()] + replacement + xml[match.end() :]
    return xml.encode("utf-8")


def _remove_calc_chain_relationship(xml_bytes: bytes) -> bytes:
    xml = xml_bytes.decode("utf-8")
    xml = re.sub(
        r'<Relationship\b[^>]*Type="http://schemas\.openxmlformats\.org/officeDocument/2006/relationships/calcChain"[^>]*/>',
        "",
        xml,
    )
    return xml.encode("utf-8")


def _remove_calc_chain_content_type(xml_bytes: bytes) -> bytes:
    xml = xml_bytes.decode("utf-8")
    xml = re.sub(r'<Override\b[^>]*PartName="/xl/calcChain\.xml"[^>]*/>', "", xml)
    return xml.encode("utf-8")


def _write_title(ws: Any, title: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center")


def _write_inputs_sheet(workbook: Workbook, inputs_df: pd.DataFrame, params: ProjectParams) -> None:
    ws = workbook.create_sheet("Исходные данные")
    _write_title(ws, "Исходные данные для групповой переоценки", len(INPUT_COLUMNS))
    params_rows = [
        ("Курс RUB за 1 CNY", params.exchange_rate_rub_per_cny),
        ("Наценка на монтаж", params.installation_markup),
        ("Валютный резерв", params.currency_reserve_share),
        ("Курс переноса RUB за 1 CNY", params.transfer_exchange_rate_rub_per_cny),
        ("Метод переноса", params.transfer_method.value),
        ("Доля переноса", params.transfer_share),
        ("Фиксированный перенос на 1 остановку, RUB", params.fixed_transfer_per_stop_rub),
    ]
    for row_idx, (label, value) in enumerate(params_rows, start=3):
        ws.cell(row_idx, 1, label)
        ws.cell(row_idx, 2, value)
    ws.cell(4, 2).number_format = PERCENT_NUMBER_FORMAT
    ws.cell(5, 2).number_format = PERCENT_NUMBER_FORMAT
    ws.cell(6, 2).number_format = "0.0000"
    ws.cell(8, 2).number_format = PERCENT_NUMBER_FORMAT
    ws.cell(9, 2).number_format = RUB_NUMBER_FORMAT

    start_row = 12
    _write_dataframe(ws, inputs_df, start_row)
    _format_table(ws, start_row, start_row + len(inputs_df), len(INPUT_COLUMNS))
    for col_idx in (3,):
        _format_column(ws, start_row + 1, start_row + len(inputs_df), col_idx, CNY_NUMBER_FORMAT)
    _format_column(ws, start_row + 1, start_row + len(inputs_df), 5, RUB_NUMBER_FORMAT)
    ws.freeze_panes = "A13"
    _set_widths(ws)


def _write_results_sheet(workbook: Workbook, results_df: pd.DataFrame) -> None:
    ws = workbook.create_sheet("Результат")
    _write_title(ws, "Результат групповой переоценки", len(RESULT_COLUMNS))
    start_row = 3
    _write_dataframe(ws, results_df, start_row)
    total_row = start_row + len(results_df) + 1
    ws.cell(total_row, 2, "Итого")
    for col_idx in [3, 6, 9, 10, 11, 12]:
        letter = get_column_letter(col_idx)
        ws.cell(total_row, col_idx, f"=SUM({letter}{start_row + 1}:{letter}{total_row - 1})")
    _format_table(ws, start_row, total_row, len(RESULT_COLUMNS))
    for col_idx in [3, 10, 11]:
        _format_column(ws, start_row + 1, total_row, col_idx, CNY_NUMBER_FORMAT)
    for col_idx in [5, 6, 7, 8, 9, 12, 13]:
        _format_column(ws, start_row + 1, total_row, col_idx, RUB_NUMBER_FORMAT)
    ws.freeze_panes = "A4"
    _set_widths(ws)


def _write_summary_sheet(workbook: Workbook, summary_df: pd.DataFrame) -> None:
    ws = workbook.create_sheet("Сводка")
    _write_title(ws, "Сводные показатели", 3)
    _write_dataframe(ws, summary_df, 3)
    _format_table(ws, 3, 3 + len(summary_df), 3)
    for row_idx in range(4, 4 + len(summary_df)):
        currency = ws.cell(row_idx, 3).value
        if currency == "CNY":
            ws.cell(row_idx, 2).number_format = CNY_NUMBER_FORMAT
        elif currency == "RUB":
            ws.cell(row_idx, 2).number_format = RUB_NUMBER_FORMAT
        elif currency == "RATE":
            ws.cell(row_idx, 2).number_format = "0.0000"
        elif currency == "PERCENT":
            ws.cell(row_idx, 2).number_format = PERCENT_NUMBER_FORMAT
    _set_widths(ws)


def _write_dataframe(ws: Any, df: pd.DataFrame, start_row: int) -> None:
    for col_idx, column in enumerate(df.columns, start=1):
        ws.cell(start_row, col_idx, column)
    for row_offset, row in enumerate(df.itertuples(index=False), start=1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(start_row + row_offset, col_idx, None if pd.isna(value) else value)


def _format_table(ws: Any, header_row: int, last_row: int, last_col: int) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=header_row, max_row=last_row, min_col=1, max_col=last_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == header_row:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            elif cell.row == last_row and ws.cell(last_row, 2).value == "Итого":
                cell.font = Font(bold=True)
                cell.fill = total_fill
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"


def _format_column(ws: Any, first_row: int, last_row: int, col_idx: int, number_format: str) -> None:
    for row_idx in range(first_row, last_row + 1):
        ws.cell(row_idx, col_idx).number_format = number_format


def _set_widths(ws: Any) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 34)
